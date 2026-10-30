"""
Script de Generación Automatizada de Informes de Prácticas Preprofesionales (UCV)
Para el Proyecto: "MÓVIL CON IA GENERATIVA PARA PREVENCIÓN DE ENFERMEDADES NO TRANSMISIBLES"
Autor: Dany Burgos Hurtado
Genera:
1. Avance 1 (Entregable 1)
2. Avance 2 (Entregable 2)
3. Informe Completo Final (Entregable 3)
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

TEMPLATE_PATH = "PPP1-C1-PRA06a-INFORME_PRACTICAS--APELLIDOS NOMBRES-.docx"
LOGO_PATH = "scratch/docx_media/word/media/image1.png"

# Paleta institucional UCV / Profesional
COLOR_PRIMARY = RGBColor(12, 44, 96)       # Azul UCV oscuro
COLOR_SECONDARY = RGBColor(194, 24, 91)    # Carmesí / Acento
COLOR_TEXT = RGBColor(33, 37, 41)          # Texto principal oscuro
COLOR_MUTED = RGBColor(108, 117, 125)      # Gris secundario

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_header_paragraph(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.bold = True
    
    if level == 1:
        run.font.size = Pt(13)
        run.font.color.rgb = COLOR_PRIMARY
    elif level == 2:
        run.font.size = Pt(12)
        run.font.color.rgb = COLOR_PRIMARY
    elif level == 3:
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(40, 60, 90)
    else:
        run.font.size = Pt(10.5)
        run.font.color.rgb = COLOR_TEXT
    return p

def add_body_paragraph(doc, text, bold_prefix=None, space_after=6, italic=False, justify=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Arial'
        r_pre.font.size = Pt(10.5)
        r_pre.bold = True
        r_pre.font.color.rgb = COLOR_TEXT
        
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(10.5)
    run.italic = italic
    run.font.color.rgb = COLOR_TEXT
    return p

def add_bullet_paragraph(doc, text, bold_title=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    r_bullet = p.add_run("• ")
    r_bullet.font.name = 'Arial'
    r_bullet.font.size = Pt(10)
    r_bullet.bold = True
    r_bullet.font.color.rgb = COLOR_PRIMARY
    
    if bold_title:
        r_title = p.add_run(bold_title + ": ")
        r_title.font.name = 'Arial'
        r_title.font.size = Pt(10)
        r_title.bold = True
        r_title.font.color.rgb = COLOR_TEXT
        
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(10)
    run.font.color.rgb = COLOR_TEXT
    return p

def add_styled_table(doc, headers, data, col_widths=None):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Header Row
    hdr_cells = table.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].text = h_text
        set_cell_background(hdr_cells[i], "0C2C60")
        set_cell_margins(hdr_cells[i], top=140, bottom=140, left=180, right=180)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = 'Arial'
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            
    # Data Rows
    for row_idx, row_data in enumerate(data):
        row_cells = table.rows[row_idx + 1].cells
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].text = str(cell_value)
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], top=100, bottom=100, left=140, right=140)
            p = row_cells[col_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx > 0 else WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(9)
                r.font.color.rgb = COLOR_TEXT
                
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table

def add_figure_caption(doc, fig_num, title, source="Elaboración propia."):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    
    r1 = p.add_run(f"Figura {fig_num}: ")
    r1.font.name = 'Arial'
    r1.font.size = Pt(9.5)
    r1.bold = True
    r1.font.color.rgb = COLOR_PRIMARY
    
    r2 = p.add_run(title)
    r2.font.name = 'Arial'
    r2.font.size = Pt(9.5)
    r2.italic = True
    r2.font.color.rgb = COLOR_TEXT
    
    p_src = doc.add_paragraph()
    p_src.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_src.paragraph_format.space_after = Pt(10)
    r3 = p_src.add_run(f"Fuente: {source}")
    r3.font.name = 'Arial'
    r3.font.size = Pt(8.5)
    r3.italic = True
    r3.font.color.rgb = COLOR_MUTED

def init_document_sections(doc):
    """
    Limpia los pies de página de la carátula (sección 0) y de las páginas preliminares
    (sección 1: resumen, índice) para que permanezcan normales y sin numeración.
    """
    if len(doc.sections) > 0:
        s0 = doc.sections[0]
        s0.footer.is_linked_to_previous = False
        for p in list(s0.footer.paragraphs):
            p._p.getparent().remove(p._p)
        s0.footer.add_paragraph()

    if len(doc.sections) > 1:
        s1 = doc.sections[1]
        s1.footer.is_linked_to_previous = False
        for p in list(s1.footer.paragraphs):
            p._p.getparent().remove(p._p)
        s1.footer.add_paragraph()

def start_body_section(doc):
    """
    Crea la sección del cuerpo principal inmediatamente después del índice
    (a partir del Capítulo I: Generalidades).
    La numeración de páginas comienza en 1 y se muestra en el pie de página
    inferior derecho con formato arábigo correlativo.
    """
    from docx.enum.section import WD_SECTION
    s_body = doc.add_section(WD_SECTION.NEW_PAGE)
    s_body.footer.is_linked_to_previous = False
    
    # Configurar para que la numeración inicie en 1 a partir del Capítulo I
    existing_pgNum = s_body._sectPr.find(qn('w:pgNumType'))
    if existing_pgNum is not None:
        existing_pgNum.set(qn('w:start'), '1')
    else:
        s_body._sectPr.append(parse_xml(r'<w:pgNumType %s w:start="1"/>' % nsdecls('w')))
        
    footer = s_body.footer
    for p in list(footer.paragraphs):
        p._p.getparent().remove(p._p)
        
    p_ftr = footer.add_paragraph()
    p_ftr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_ftr.paragraph_format.space_before = Pt(4)
    p_ftr.paragraph_format.space_after = Pt(0)
    
    # Campo dinámico de página OpenXML (compatible con Word Desktop, Web y LibreOffice)
    fld_xml = parse_xml(r'''
        <w:fldSimple %s w:instr="PAGE">
            <w:r>
                <w:rPr>
                    <w:rFonts w:ascii="Arial" w:hAnsi="Arial"/>
                    <w:sz w:val="20"/>
                    <w:color w:val="6C757D"/>
                </w:rPr>
                <w:t>1</w:t>
            </w:r>
        </w:fldSimple>
    ''' % nsdecls('w'))
    
    p_ftr._p.append(fld_xml)
    return s_body

print("Módulo de estilos y utilidades docx cargado con éxito.")
